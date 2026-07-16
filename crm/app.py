import io
import sqlite3
from datetime import date
from pathlib import Path

from flask import Flask, g, redirect, render_template, request, send_file, url_for
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent
DB_PATH = BASE_DIR / "customers.db"

SALE_STATUSES = ["見込み", "受注", "請求済み", "入金済み"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

app = Flask(__name__)


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    with app.app_context():
        db = get_db()
        with open(BASE_DIR / "schema.sql") as f:
            db.executescript(f.read())
        db.commit()


@app.route("/")
def index():
    db = get_db()
    customer_count = db.execute("SELECT COUNT(*) FROM customers").fetchone()[0]
    sale_count = db.execute("SELECT COUNT(*) FROM sales").fetchone()[0]
    total_amount = db.execute("SELECT COALESCE(SUM(amount), 0) FROM sales").fetchone()[0]
    month_amount = db.execute(
        "SELECT COALESCE(SUM(amount), 0) FROM sales "
        "WHERE strftime('%Y-%m', sale_date) = strftime('%Y-%m', 'now', 'localtime')"
    ).fetchone()[0]
    recent_sales = db.execute(
        "SELECT sales.*, customers.company_name FROM sales "
        "JOIN customers ON customers.id = sales.customer_id "
        "ORDER BY sales.sale_date DESC, sales.id DESC LIMIT 5"
    ).fetchall()
    monthly_sales = _monthly_chart_data(db)
    return render_template(
        "dashboard.html",
        customer_count=customer_count,
        sale_count=sale_count,
        total_amount=total_amount,
        month_amount=month_amount,
        recent_sales=recent_sales,
        monthly_sales=monthly_sales,
    )


@app.route("/customers")
def list_customers():
    db = get_db()
    customers = db.execute(
        "SELECT customers.*, COALESCE(SUM(sales.amount), 0) AS total_amount "
        "FROM customers LEFT JOIN sales ON sales.customer_id = customers.id "
        "GROUP BY customers.id ORDER BY customers.id DESC"
    ).fetchall()
    return render_template("customer_list.html", customers=customers)


@app.route("/customers/new", methods=["GET", "POST"])
def new_customer():
    if request.method == "POST":
        db = get_db()
        db.execute(
            "INSERT INTO customers (company_name, contact_name, email, phone, notes) "
            "VALUES (?, ?, ?, ?, ?)",
            (
                request.form["company_name"],
                request.form["contact_name"],
                request.form["email"],
                request.form["phone"],
                request.form["notes"],
            ),
        )
        db.commit()
        return redirect(url_for("list_customers"))
    return render_template("customer_form.html", customer=None)


@app.route("/customers/edit/<int:customer_id>", methods=["GET", "POST"])
def edit_customer(customer_id):
    db = get_db()
    if request.method == "POST":
        db.execute(
            "UPDATE customers SET company_name = ?, contact_name = ?, email = ?, "
            "phone = ?, notes = ? WHERE id = ?",
            (
                request.form["company_name"],
                request.form["contact_name"],
                request.form["email"],
                request.form["phone"],
                request.form["notes"],
                customer_id,
            ),
        )
        db.commit()
        return redirect(url_for("list_customers"))
    customer = db.execute(
        "SELECT * FROM customers WHERE id = ?", (customer_id,)
    ).fetchone()
    return render_template("customer_form.html", customer=customer)


@app.route("/customers/delete/<int:customer_id>", methods=["POST"])
def delete_customer(customer_id):
    db = get_db()
    db.execute("DELETE FROM customers WHERE id = ?", (customer_id,))
    db.commit()
    return redirect(url_for("list_customers"))


def _yearly_chart_data(rows):
    yearly_totals = {int(r["year"]): r["total"] for r in rows if r["year"]}
    if not yearly_totals:
        return []
    max_total = max(yearly_totals.values()) or 1
    return [
        {
            "year": year,
            "total": yearly_totals.get(year, 0),
            "pct": round(yearly_totals.get(year, 0) / max_total * 100),
        }
        for year in range(min(yearly_totals), max(yearly_totals) + 1)
    ]


def _monthly_chart_data(db, customer_id=None, months=12):
    query = "SELECT strftime('%Y-%m', sale_date) AS ym, SUM(amount) AS total FROM sales"
    params = []
    if customer_id is not None:
        query += " WHERE customer_id = ?"
        params.append(customer_id)
    query += " GROUP BY ym"
    totals = {r["ym"]: r["total"] for r in db.execute(query, params).fetchall()}

    today = date.today()
    keys = []
    year, month = today.year, today.month
    for i in range(months - 1, -1, -1):
        m = month - i
        y = year
        while m <= 0:
            m += 12
            y -= 1
        keys.append((y, m))

    max_total = max([totals.get(f"{y:04d}-{m:02d}", 0) for y, m in keys] + [0]) or 1
    return [
        {
            "label": f"{y % 100:02d}/{m}",
            "total": totals.get(f"{y:04d}-{m:02d}", 0),
            "pct": round(totals.get(f"{y:04d}-{m:02d}", 0) / max_total * 100),
        }
        for y, m in keys
    ]


def _sales_workbook(rows, sheet_title, include_customer):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    headers = ["日付", "案件名", "金額", "ステータス", "備考"]
    if include_customer:
        headers.insert(1, "顧客")
    ws.append(headers)
    total = 0
    for r in rows:
        total += r["amount"]
        row = [r["sale_date"], r["item_name"], r["amount"], r["status"], r["memo"]]
        if include_customer:
            row.insert(1, r["company_name"])
        ws.append(row)
    ws.append([])
    total_row = [""] * (len(headers) - 2) + ["合計", total]
    ws.append(total_row)
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    return wb


def _send_workbook(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename, mimetype=XLSX_MIME)


@app.route("/customers/<int:customer_id>")
def customer_detail(customer_id):
    db = get_db()
    customer = db.execute(
        "SELECT * FROM customers WHERE id = ?", (customer_id,)
    ).fetchone()
    if customer is None:
        return redirect(url_for("list_customers"))
    sales = db.execute(
        "SELECT * FROM sales WHERE customer_id = ? ORDER BY sale_date DESC, id DESC",
        (customer_id,),
    ).fetchall()
    total_amount = sum(s["amount"] for s in sales)

    chart_view = request.args.get("chart", "yearly")
    if chart_view not in ("yearly", "monthly"):
        chart_view = "yearly"
    if chart_view == "monthly":
        chart_data = _monthly_chart_data(db, customer_id=customer_id)
    else:
        yearly_rows = db.execute(
            "SELECT strftime('%Y', sale_date) AS year, SUM(amount) AS total "
            "FROM sales WHERE customer_id = ? GROUP BY year ORDER BY year",
            (customer_id,),
        ).fetchall()
        chart_data = _yearly_chart_data(yearly_rows)

    return render_template(
        "customer_detail.html",
        customer=customer,
        sales=sales,
        total_amount=total_amount,
        chart_data=chart_data,
        chart_view=chart_view,
    )


@app.route("/customers/<int:customer_id>/export")
def export_customer_sales(customer_id):
    db = get_db()
    customer = db.execute(
        "SELECT * FROM customers WHERE id = ?", (customer_id,)
    ).fetchone()
    if customer is None:
        return redirect(url_for("list_customers"))
    sales = db.execute(
        "SELECT * FROM sales WHERE customer_id = ? ORDER BY sale_date DESC, id DESC",
        (customer_id,),
    ).fetchall()
    wb = _sales_workbook(sales, sheet_title="売上履歴", include_customer=False)
    return _send_workbook(wb, f"{customer['company_name']}_売上.xlsx")


@app.route("/sales")
def list_sales():
    db = get_db()
    sales = db.execute(
        "SELECT sales.*, customers.company_name FROM sales "
        "JOIN customers ON customers.id = sales.customer_id "
        "ORDER BY sales.sale_date DESC, sales.id DESC"
    ).fetchall()
    total_amount = sum(s["amount"] for s in sales)
    return render_template("sale_list.html", sales=sales, total_amount=total_amount)


@app.route("/sales/export")
def export_sales():
    db = get_db()
    sales = db.execute(
        "SELECT sales.*, customers.company_name FROM sales "
        "JOIN customers ON customers.id = sales.customer_id "
        "ORDER BY sales.sale_date DESC, sales.id DESC"
    ).fetchall()
    wb = _sales_workbook(sales, sheet_title="売上一覧", include_customer=True)
    return _send_workbook(wb, "売上一覧.xlsx")


@app.route("/sales/new", methods=["GET", "POST"])
def new_sale():
    db = get_db()
    if request.method == "POST":
        db.execute(
            "INSERT INTO sales (customer_id, item_name, amount, status, sale_date, memo) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (
                request.form["customer_id"],
                request.form["item_name"],
                request.form["amount"],
                request.form["status"],
                request.form["sale_date"],
                request.form["memo"],
            ),
        )
        db.commit()
        return redirect(url_for("list_sales"))
    customers = db.execute("SELECT * FROM customers ORDER BY company_name").fetchall()
    selected_customer_id = request.args.get("customer_id", type=int)
    return render_template(
        "sale_form.html",
        sale=None,
        customers=customers,
        statuses=SALE_STATUSES,
        selected_customer_id=selected_customer_id,
    )


@app.route("/sales/edit/<int:sale_id>", methods=["GET", "POST"])
def edit_sale(sale_id):
    db = get_db()
    if request.method == "POST":
        db.execute(
            "UPDATE sales SET customer_id = ?, item_name = ?, amount = ?, "
            "status = ?, sale_date = ?, memo = ? WHERE id = ?",
            (
                request.form["customer_id"],
                request.form["item_name"],
                request.form["amount"],
                request.form["status"],
                request.form["sale_date"],
                request.form["memo"],
                sale_id,
            ),
        )
        db.commit()
        return redirect(url_for("list_sales"))
    sale = db.execute("SELECT * FROM sales WHERE id = ?", (sale_id,)).fetchone()
    customers = db.execute("SELECT * FROM customers ORDER BY company_name").fetchall()
    return render_template(
        "sale_form.html",
        sale=sale,
        customers=customers,
        statuses=SALE_STATUSES,
        selected_customer_id=sale["customer_id"] if sale else None,
    )


@app.route("/sales/delete/<int:sale_id>", methods=["POST"])
def delete_sale(sale_id):
    db = get_db()
    db.execute("DELETE FROM sales WHERE id = ?", (sale_id,))
    db.commit()
    return redirect(url_for("list_sales"))


if __name__ == "__main__":
    init_db()
    app.run(debug=True, port=5000)
